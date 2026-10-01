import os
import glob
from pathlib import Path
from config import DATA_DIR, GEMINI_API_KEY
from google import genai
from docx import Document

# Configura Gemini
if GEMINI_API_KEY:
    client = genai.Client(api_key=GEMINI_API_KEY)
else:
    client = None

def atualizar_arquivos_locais(mudancas):
    if not client or not mudancas:
        return
    
    # Filtra apenas as mudanças reais de texto para passar para a IA
    texto_mudancas = "\n".join([m for m in mudancas if "NOVA PÁGINA" not in m and "PÁGINA REMOVIDA" not in m and "Nenhuma alteração" not in m])
    
    if not texto_mudancas.strip():
        return

    print("\n[Atualizador IA] Verificando se os arquivos locais (TXT, DOCX) precisam ser atualizados com base nas mudanças do site...")
    
    # Processa TXT
    txt_files = glob.glob(str(DATA_DIR / "*.txt"))
    for txt_file in txt_files:
        if "base_conhecimento" in txt_file:
            continue
        try:
            with open(txt_file, "r", encoding="utf-8") as f:
                conteudo = f.read()
            
            novo_conteudo = reescrever_com_ia(conteudo, texto_mudancas)
            if novo_conteudo and novo_conteudo != conteudo:
                with open(txt_file, "w", encoding="utf-8") as f:
                    f.write(novo_conteudo)
                print(f" -> Arquivo TXT atualizado: {Path(txt_file).name}")
            import time
            time.sleep(4) # Evita limite de 15 RPM da API Gratuita
        except Exception as e:
            print(f"Erro ao processar TXT {txt_file}: {e}")

    # Processa DOCX
    docx_files = glob.glob(str(DATA_DIR / "*.docx"))
    for docx_file in docx_files:
        try:
            doc = Document(docx_file)
            conteudo_docx = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
            if not conteudo_docx.strip():
                continue
                
            novo_conteudo = reescrever_com_ia(conteudo_docx, texto_mudancas)
            if novo_conteudo and novo_conteudo != conteudo_docx:
                # Cria um novo docx com o texto atualizado (simplificado)
                novo_doc = Document()
                for linha in novo_conteudo.split('\n'):
                    if linha.strip():
                        novo_doc.add_paragraph(linha)
                novo_doc.save(docx_file)
                print(f" -> Arquivo DOCX atualizado: {Path(docx_file).name}")
            import time
            time.sleep(4) # Evita limite de 15 RPM da API Gratuita
        except Exception as e:
            print(f"Erro ao processar DOCX {docx_file}: {e}")
            
    print("[Atualizador IA] A atualização de arquivos PDF não é suportada por modificar a estrutura visual original. Recomendamos atualizar o PDF a partir do Word.")

def reescrever_com_ia(texto_original, mudancas):
    prompt = f"""Você é um assistente que atualiza documentos baseados em novas regras.
Abaixo estão as mudanças detectadas no site oficial:
{mudancas}

Aqui está o conteúdo do documento local:
---
{texto_original}
---

Sua tarefa:
Analise as mudanças do site. Se as mudanças conflitarem ou atualizarem regras presentes no documento local (por exemplo, limite de saques mudou de 3 para 4), reescreva o documento local atualizando os valores/regras. 
Se o documento não possuir relação com as mudanças, retorne exatamente o documento original.
Mantenha o formato de texto. Retorne APENAS o texto atualizado, sem comentários adicionais."""
    
    import time
    for tentativa in range(3):
        try:
            response = client.models.generate_content(
                model='gemini-1.5-flash',
                contents=prompt
            )
            return response.text.strip()
        except Exception as e:
            if ("503" in str(e) or "429" in str(e)) and tentativa < 2:
                print(f"[Atualizador IA] Servidor ocupado (tentativa {tentativa+1}/3). Aguardando 5 segundos...")
                time.sleep(5)
                continue
            print(f"Erro na API do Gemini: {e}")
            return texto_original
